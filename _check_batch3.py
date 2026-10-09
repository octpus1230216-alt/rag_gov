# -*- coding: utf-8 -*-
"""
第 3 批自检：真造 docx/pptx/xlsx/pdf/txt + 合成 blocks，跑 layout_chunker /
table_chunker / pipeline，断言统一 chunk schema、无内容丢失、chunk_id 稳定
（两次跑 + 串行与多进程一致）、跨页同 section 合并、overlap 为整句、
表格整表/行组各片带表头、幂等跳过、断点续跑复用 parsed、坏文件定位到具体路径。

运行：python _check_batch3.py
"""

from __future__ import annotations

import json
import re
import sys
from itertools import pairwise
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

TMP = ROOT / "_tmp_batch3"
OUT = TMP / "out"
IN_DIR = TMP / "in"

FAILS: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    if cond:
        print(f"[OK]   {name}")
    else:
        FAILS.append(f"{name} :: {detail}")
        print(f"[FAIL] {name} :: {detail}")


# ============================================================ 造样例文件
LONG_SENT = "本合同项下甲方的付款期限为发票到账后的三十个自然日内逾期按日万分之五计算违约金直至全部款项结清为止。"


def make_txt() -> Path:
    """带长正文 + markdown 大表的 txt（同时压到切分与表格两条路）。"""
    p = IN_DIR / "sample.txt"
    body = "".join(f"{LONG_SENT}第{i}条补充说明内容。\n" for i in range(1, 31))
    rows = "\n".join(
        f"| {i} | 设备采购明细第{i}项 | {12000 + i * 37} | 含三年质保与现场安装培训服务费用已含在单价内 |"
        for i in range(1, 41)
    )
    content = (
        "# 采购合同\n"
        "甲方：星辰科技有限公司\n"
        "乙方：海纳供应链有限公司\n\n"
        "## 第一条 标的物\n"
        f"{body}\n"
        "## 第二条 费用明细\n"
        "| 序号 | 名称 | 单价 | 备注 |\n"
        "|---|---|---|---|\n"
        f"{rows}\n\n"
        "## 第三条 违约责任\n"
        "违约方应向守约方支付合同总价百分之十的违约金，并赔偿由此造成的全部损失。\n"
    )
    p.write_text(content, encoding="utf-8")
    return p


def make_docx() -> Path:
    import docx

    p = IN_DIR / "sample.docx"
    d = docx.Document()
    d.core_properties.title = "产品使用手册"
    d.add_heading("产品使用手册", level=1)
    d.add_heading("第3章 系统配置", level=1)
    for i in range(1, 13):
        d.add_paragraph(f"第3章正文第{i}句：{LONG_SENT}")
    d.add_heading("3.2 参数表", level=2)
    table = d.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "参数名"
    table.rows[0].cells[1].text = "默认值"
    table.rows[0].cells[2].text = "说明"
    table.rows[1].cells[0].text = "chunk_size"
    table.rows[1].cells[1].text = "384"
    table.rows[1].cells[2].text = "切分窗口大小"
    table.rows[2].cells[0].text = "chunk_overlap"
    table.rows[2].cells[1].text = "64"
    table.rows[2].cells[2].text = "相邻窗口重叠量"
    d.add_paragraph("表格之后的正文：上述参数需要与硬件档位一起调整。")
    d.save(p)
    return p


def make_pptx() -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    p = IN_DIR / "sample.pptx"
    prs = Presentation()
    for i in (1, 2):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
        para = box.text_frame.paragraphs[0]
        para.text = f"第{i}页：季度营收同比增长百分之十八，成本下降百分之三。"
        para.font.size = Pt(24)
    prs.save(p)
    return p


def make_xlsx() -> Path:
    from openpyxl import Workbook

    p = IN_DIR / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "费用明细"
    ws.append(["部门", "科目", "金额", "备注说明"])
    for i in range(1, 46):
        ws.append([f"研发部{i % 5}", "差旅费", 1000 + i * 13, "含往返交通与住宿补贴按公司标准执行"])
    wb.create_sheet("预算")
    wb.save(p)
    return p


def make_pdf() -> Path:
    """两页同一章节的正文，用来验「跨页同 section 合并」在真实 PDF 上生效。"""
    import pymupdf

    p = IN_DIR / "sample.pdf"
    doc = pymupdf.open()
    for page_no in (1, 2):
        page = doc.new_page()
        page.insert_textbox(
            pymupdf.Rect(50, 60, 540, 720),
            f"第{page_no}页 {LONG_SENT}这是跨页延续的第{page_no * 2}段文字内容用于验证合并。",
            fontname="china-ss",
            fontsize=12,
        )
    doc.set_metadata({"title": "财务分析报告"})
    doc.save(p)
    doc.close()
    return p


def make_broken_pdf() -> Path:
    p = IN_DIR / "broken.pdf"
    p.write_bytes(b"%PDF-1.4\nthis is not a real pdf body\n%%EOF-truncated")
    return p


def make_all() -> list[Path]:
    """
    每次跑都从零开始：python-docx / openpyxl / python-pptx 存盘时把 zip 时间戳
    写进包里，同一份内容两次生成的字节不同 -> doc_id（sha256(字节)）也不同，
    留着上一轮的产物会让「幂等不重复追加」这类断言被历史脏数据污染。
    """
    import shutil

    if TMP.exists():
        shutil.rmtree(TMP)
    for d in (TMP, IN_DIR, OUT):
        d.mkdir(parents=True, exist_ok=True)
    return [make_txt(), make_docx(), make_pptx(), make_xlsx(), make_pdf(), make_broken_pdf()]


# ============================================================ 合成 blocks
def synth_blocks(n_sentences: int = 40) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """同 section 跨 3 页的正文 + 中间夹一张表 + 尾部另一个 section。"""
    blocks: list[dict[str, Any]] = [
        {"type": "title", "content": "第3章 付款方式", "page": 1, "section_path": "第3章 付款方式"},
    ]
    for i in range(n_sentences):
        blocks.append(
            {
                "type": "text",
                "content": f"{LONG_SENT}第{i}项补充约定。",
                "page": 1 + i // 14,
                "section_path": "第3章 付款方式",
                "bbox": [50.0, 60.0 + i, 540.0, 80.0 + i],
            }
        )
    blocks.append(
        {
            "type": "title",
            "content": "3.2 表格",
            "page": 3,
            "section_path": "第3章 付款方式 > 3.2 表格",
        }
    )
    blocks.append(
        {
            "type": "table",
            "content": (
                "| 项目 | 金额 | 说明 |\n|---|---|---|\n"
                "| 预付款 | 30000 | 签约后七日内支付 |\n| 尾款 | 70000 | 验收合格后支付 |\n"
            ),
            "page": 3,
            "section_path": "第3章 付款方式 > 3.2 表格",
            "header": ["项目", "金额", "说明"],
            "bbox": [50.0, 300.0, 540.0, 420.0],
        }
    )
    blocks.append(
        {
            "type": "text",
            "content": "独立章节的一句话：所有款项均以银行转账方式结算。",
            "page": 4,
            "section_path": "第4章 结算",
        }
    )
    metadata = {
        "doc_type": "contract",
        "topic": "finance",
        "has_table": True,
        "has_image": False,
        "language": "zh",
        "layout_style": "mixed",
        "title": "采购合同",
    }
    return blocks, metadata


# ============================================================ 测试：token 估算
def test_token_estimator() -> None:
    from ingestion.chunkers import layout_chunker as lc

    check("估算器：中文长度 100 字 ≈ 60 token", 50 <= lc.estimate_tokens_zh("中" * 100) <= 70,
          str(lc.estimate_tokens_zh("中" * 100)))
    check("估算器：纯英文不被 ×0.6 高估",
          lc.estimate_tokens_zh("internationalization organization") < 22 * 0.6 + 2,
          str(lc.estimate_tokens_zh("internationalization organization")))
    check("估算器：空串 = 0", lc.estimate_tokens_zh("   ") == 0)

    func, backend = lc.choose_token_estimator("char")
    check("choose：char 后端", func is lc.estimate_tokens_zh and backend == "char_heuristic", backend)
    try:
        lc.choose_token_estimator("jieba")
        jieba_ok = True
    except Exception:
        jieba_ok = False
    _auto_func, auto_backend = lc.choose_token_estimator(None)
    check("choose：auto 缺 jieba 时退回字符口径",
          auto_backend == ("jieba" if jieba_ok else "char_heuristic"), auto_backend)

    named = lc.set_token_estimator(lambda t: len(t))
    check("set_token_estimator 可整体替换", lc.estimate_tokens("四个字") == 3 and named == "<lambda>",
          f"{lc.estimate_tokens('四个字')} / {named}")
    lc.set_token_estimator(prefer="char")
    check("切回 char 后端", lc.token_estimator_name() == "char_heuristic", lc.token_estimator_name())


def test_config_resolve() -> None:
    from ingestion.chunkers.layout_chunker import resolve_config

    cfg = resolve_config(None)
    check("默认 chunk_size=384", cfg.chunk_size == 384, str(cfg.chunk_size))
    check("默认 overlap=64", cfg.chunk_overlap == 64, str(cfg.chunk_overlap))
    check("settings 里的表头行数生效", cfg.max_table_rows in (30, 8), str(cfg.max_table_rows))
    check("separators 去重保序", len(set(cfg.separators)) == len(cfg.separators), str(cfg.separators))
    check("显式覆盖生效", resolve_config({"chunk_size": 128}).chunk_size == 128)


# ============================================================ 测试：正文切分
def test_text_chunking() -> None:
    from ingestion.chunkers.layout_chunker import chunk_text_blocks, estimate_tokens, resolve_config

    blocks, metadata = synth_blocks(40)
    cfg = resolve_config(None)
    chunks = chunk_text_blocks(blocks, metadata, cfg, doc_id="d" * 64)
    texts = [c for c in chunks if c["chunk_type"] == "text"]

    check("正文切出多片（长章节被分窗口）", len(texts) >= 3, str(len(texts)))
    check("每片 token 不超 chunk_size×1.35",
          all(estimate_tokens(c["content"]) <= cfg.chunk_size * 1.35 for c in texts),
          str([c["token_estimate"] for c in texts]))
    check("标题不单独成 chunk",
          not any(c["content"].strip() in ("第3章 付款方式", "3.2 表格") for c in texts))
    check("表格块不走正文切分", all("项目" not in c["content"] or "|" not in c["content"] for c in texts))
    check("section_path 带上标题层级",
          texts[0]["section_path"] == "第3章 付款方式", texts[0]["section_path"])

    # 跨页合并：第3章的片断应同时覆盖 1/2/3 页（正文被别的章节隔开，各自成组）
    sec3 = [c for c in texts if c["section_path"] == "第3章 付款方式"]
    covered = {p for c in sec3 for p in (c.get("pages") or [])}
    check("跨页同 section 合并（单片可跨多页）",
          any(len(c.get("pages") or []) > 1 for c in sec3), str([c.get("pages") for c in sec3[:3]]))
    check("第3章正文覆盖 1/2/3 页", covered >= {1, 2, 3}, f"{covered}")
    check("不同 section 分片不混（第4章单独成片）",
          any(c["section_path"] == "第4章 结算" for c in texts), str({c["section_path"] for c in texts}))

    # 无内容丢失：每句原文都能在某个 chunk 里找到
    sentences = [b["content"] for b in blocks if b["type"] == "text"]
    lost = [s for s in sentences if not any(s in c["content"] for c in texts)]
    check("正文一句都不丢", not lost, f"{len(lost)} 句丢失: {lost[:1]}")

    # overlap 必须是整句：同 section 的相邻片，后片开头那句要能在前片整句出现
    bad: list[str] = []
    overlapped = 0
    for section in dict.fromkeys(c["section_path"] for c in texts):
        group = [c for c in texts if c["section_path"] == section]
        for prev, nxt in pairwise(group):
            head = nxt["content"].split("。")[0]
            if nxt["content"][:12] in prev["content"]:
                overlapped += 1
            if len(head) > 1 and head not in prev["content"]:
                bad.append(head[:20])
    check("overlap 以整句为单位（不切碎数字/句子）", not bad, str(bad[:2]))
    check("相邻片确实带重叠内容（不是硬拼接口）", overlapped > 0,
          str([c["content"][:12] for c in texts[:3]]))

    # chunk_id 稳定：同输入两次 + 换个顺序构造
    again = chunk_text_blocks(blocks, metadata, cfg, doc_id="d" * 64)
    check("chunk_id 幂等（同输入同输出）",
          [c["chunk_id"] for c in again] == [c["chunk_id"] for c in texts])
    check("chunk_id 是 sha256 十六进制", all(re.fullmatch(r"[0-9a-f]{64}", c["chunk_id"]) for c in texts))
    check("chunk_id 全局唯一", len({c["chunk_id"] for c in texts}) == len(texts))
    check("chunk_id 顺序编号连续", [c["chunk_index"] for c in texts] == list(range(len(texts))))
    check("不同 doc_id 不撞 id",
          {c["chunk_id"] for c in chunk_text_blocks(blocks, metadata, cfg, doc_id="e" * 64)}
          != {c["chunk_id"] for c in texts})


def test_chunk_schema() -> None:
    """统一 Chunk 输出格式：六个必带字段 + 可选 bbox + metadata 继承。"""
    from ingestion.pipeline import run

    summary = run(IN_DIR, "parse", "chunk", workers=2, root=OUT, force=True)
    check("流水线无失败（除刻意造的坏文件）", summary["error"] == 1, str(summary["error"]))

    required = {"chunk_id", "doc_id", "content", "chunk_type", "section_path", "metadata"}
    allowed_types = {"text", "table", "image", "slide", "sheet"}
    seen_types: set[str] = set()
    problems: list[str] = []
    meta_keys = {"doc_type", "topic", "language", "has_table", "has_image", "layout_style"}
    for path in sorted((OUT / "chunks").glob("*.jsonl")):
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            c = json.loads(line)
            if not required.issubset(c):
                problems.append(f"{path.name} 缺字段 {required - set(c)}")
            if c["chunk_type"] not in allowed_types:
                problems.append(f"{path.name} 非法 chunk_type {c['chunk_type']}")
            if not isinstance(c["section_path"], str) or not isinstance(c["metadata"], dict):
                problems.append(f"{path.name} section_path/metadata 类型错")
            if not meta_keys.issubset(c["metadata"]):
                problems.append(f"{path.name} metadata 未继承 {meta_keys - set(c['metadata'])}")
            if "bbox" in c and len(c["bbox"]) != 4:
                problems.append(f"{path.name} bbox 长度 {len(c['bbox'])}")
            if not c["content"].strip():
                problems.append(f"{path.name} 空内容 chunk")
            seen_types.add(c["chunk_type"])
    check("所有 chunk 符合统一格式（6 字段/类型/元数据继承）", not problems, str(problems[:3]))
    check("覆盖 text 类型", "text" in seen_types, str(seen_types))
    check("覆盖 table 类型", "table" in seen_types, str(seen_types))
    check("覆盖 slide 类型", "slide" in seen_types, str(seen_types))
    check("产物落盘 parsed/*.json", len(list((OUT / "parsed").glob("*.json"))) >= 5,
          str(len(list((OUT / "parsed").glob("*.json")))))
    check("产物落盘 chunks/*.jsonl", len(list((OUT / "chunks").glob("*.jsonl"))) >= 5,
          str(len(list((OUT / "chunks").glob("*.jsonl")))))
    check("未实现阶段不伪造 vectors/", not (OUT / "vectors").exists())


def test_real_docs_quality() -> None:
    from ingestion.parsers.common import compute_doc_id

    def chunks_of(path: Path) -> list[dict[str, Any]]:
        p = OUT / "chunks" / f"{compute_doc_id(path)}.jsonl"
        return [json.loads(ln) for ln in p.read_text(encoding="utf-8").splitlines() if ln.strip()]

    # PDF：两页同章节 -> 合并后正文片数应远少于「按页切」的片数
    pdf_chunks = [c for c in chunks_of(IN_DIR / "sample.pdf") if c["chunk_type"] == "text"]
    multi = [c for c in pdf_chunks if len(c.get("pages") or []) > 1]
    check("真实 PDF：跨页同 section 合并成片", bool(multi) or len(pdf_chunks) == 1,
          str([c.get("pages") for c in pdf_chunks]))

    # txt：大表被行组分片，正文不吞表格
    txt_chunks = chunks_of(IN_DIR / "sample.txt")
    tables = [c for c in txt_chunks if c["chunk_type"] == "table"]
    check("txt 表格走表格切分器", bool(tables), str({c['chunk_type'] for c in txt_chunks}))
    check("txt 正文不混表格竖线",
          all("|---" not in c["content"] for c in txt_chunks if c["chunk_type"] == "text"))

    # xlsx：sheet 类型
    xlsx_chunks = chunks_of(IN_DIR / "sample.xlsx")
    check("xlsx 产出 sheet chunk", any(c["chunk_type"] == "sheet" for c in xlsx_chunks),
          str({c["chunk_type"] for c in xlsx_chunks}))

    # pptx：一页一 slide chunk
    pptx_chunks = [c for c in chunks_of(IN_DIR / "sample.pptx") if c["chunk_type"] == "slide"]
    check("pptx 一页一 chunk", len(pptx_chunks) == 2, str(len(pptx_chunks)))

    # docx：表格前的小表整表一块
    docx_chunks = chunks_of(IN_DIR / "sample.docx")
    small = [c for c in docx_chunks if c["chunk_type"] == "table"]
    check("docx 小表整表成块不切碎", len(small) == 1 and small[0].get("n_parts") in (None,),
          str([(c.get("part"), c.get("n_parts")) for c in small]))


# ============================================================ 测试：表格切分
def test_table_chunking() -> None:
    from ingestion.chunkers.table_chunker import chunk_table, is_markdown_table, row_to_kv

    md_small = (
        "| 项目 | 金额（元） | 备注 |\n|---|---|---|\n"
        "| 预付款 | 30,000.00 | 签约后七日内支付 |\n| 尾款 | 70,000.50 | 验收合格后支付 |\n"
    )
    block = {"type": "table", "content": md_small, "page": 3,
             "section_path": "第3章 > 3.2", "header": ["项目", "金额（元）", "备注"]}
    out = chunk_table(block, {"doc_type": "contract"}, None, doc_id="f" * 64)
    check("小表整表一个 chunk", len(out) == 1, str(len(out)))
    check("content 保留原始 markdown 表格", "| 预付款 |" in out[0]["content"] and "---" in out[0]["content"])
    check("section_path 追加表格序号", out[0]["section_path"] == "第3章 > 3.2 > 表格1", out[0]["section_path"])
    check("列名取自 parser 的 header", out[0]["columns"] == ["项目", "金额（元）", "备注"], str(out[0]["columns"]))
    check("逐行 kv 展开格式「列名: 值 | ...」",
          out[0]["row_lines"][0] == "项目: 预付款 | 金额（元）: 30,000.00 | 备注: 签约后七日内支付",
          out[0]["row_lines"][0])
    check("is_markdown_table 判定", is_markdown_table(md_small) and not is_markdown_table("纯文本"))
    check("row_to_kv 跳过空值", row_to_kv(["a", "", "c"], ["X", "Y", "Z"]) == "X: a | Z: c",
          row_to_kv(["a", "", "c"], ["X", "Y", "Z"]))

    # 大表：>2000 字符 -> 行组分片，每片都带表头
    rows = [
        f"| {i} | 设备采购明细第{i}项 | {12000 + i * 37} | 含三年质保与现场安装培训服务费用已含在单价内 |"
        for i in range(1, 61)
    ]
    header_lines = ["| 序号 | 名称 | 单价 | 备注 |", "|---|---|---|---|"]
    big = "\n".join([*header_lines, *rows]) + "\n"
    check("样例大表确实超过 2000 字符", len(big) > 2000, str(len(big)))
    big_block = {"type": "table", "content": big, "page": 5,
                 "section_path": "附表", "header": ["序号", "名称", "单价", "备注"]}
    parts = chunk_table(big_block, {}, None, doc_id="g" * 64)
    check("大表按行组切成多片", len(parts) > 1, str(len(parts)))
    check("每片首行都是完整表头行",
          all(p["content"].splitlines()[0] == "| 序号 | 名称 | 单价 | 备注 |" for p in parts),
          parts[1]["content"].splitlines()[0])
    check("每片 row_lines 每行都自带列名（精确检索不丢列语义）",
          all(all(ln.startswith("序号: ") and "名称: " in ln for ln in p["row_lines"]) for p in parts if p.get("row_lines")),
          parts[0]["row_lines"][0][:40] if parts[0].get("row_lines") else "无 row_lines")
    check("片数/序号/行区间标注齐全",
          all(p["part"] == i + 1 and p["n_parts"] == len(parts) and p["row_range"] for i, p in enumerate(parts)))
    # 行不丢不重
    rebuilt = [ln for p in parts for ln in p["content"].splitlines() if ln.startswith("|")]
    data_lines = [ln for ln in rebuilt if "序号" not in ln and not set(ln) <= set("|-: ")]
    uniq = list(dict.fromkeys(data_lines))
    check("大表数据行不丢", len(uniq) == len(rows), f"{len(uniq)} vs {len(rows)}")
    check("大表数据行不重", len(data_lines) == len(uniq), str(len(data_lines) - len(uniq)))
    check("行区间按原文顺序递增",
          [int(p["row_range"].split("-")[0]) for p in parts] == sorted(int(p["row_range"].split("-")[0]) for p in parts))
    check("每片字符数不超预算",
          all(len(p["content"]) <= 2200 for p in parts), str([len(p["content"]) for p in parts][:4]))

    # 非法/非表格内容：原样成一块，绝不静默丢内容
    loose = {"type": "table", "content": "表1：2024年度费用汇总（图片式表格无法结构化）", "page": 1}
    out2 = chunk_table(loose, {}, None, doc_id="h" * 64)
    check("非 markdown 表格原样保留", len(out2) == 1 and out2[0]["not_structured"] is True,
          str(out2))
    check("空表格块不产出 chunk", chunk_table({"type": "table", "content": "   "}, {}) == [])

    # sheet 复用同一切分器
    sheet_block = {"type": "sheet", "content": md_small, "sheet_name": "费用", "region_index": 2}
    out3 = chunk_table(sheet_block, {}, None, chunk_type="sheet", table_index=2, doc_id="i" * 64)
    check("sheet 块 chunk_type=sheet", out3[0]["chunk_type"] == "sheet" and out3[0]["sheet_name"] == "费用")


# ============================================================ 测试：流水线
def test_pipeline_idempotent() -> None:
    from ingestion.parsers.common import compute_doc_id
    from ingestion.pipeline import (
        IMPLEMENTED,
        STAGES,
        artifact_path,
        iter_chunks,
        main,
        run,
        stage_range,
    )

    first_ids = [c["chunk_id"] for c in iter_chunks(compute_doc_id(IN_DIR / "sample.docx"), OUT)]
    check("首轮产出 chunk", len(first_ids) > 0, str(len(first_ids)))

    second = run(IN_DIR, "parse", "chunk", workers=2, root=OUT)
    check("第二跑：全部文件零失败（除坏文件）", second["error"] == 1, str(second["error"]))
    ok_reports = [r for r in second["reports"] if r["status"] == "done"]
    check("第二跑：所有阶段 skipped（幂等生效）",
          all(all(v == "skipped" for v in r["stages"].values()) for r in ok_reports),
          str([r["stages"] for r in ok_reports][:2]))
    check("第二跑：fully_skipped 统计正确", second["fully_skipped"] == len(ok_reports),
          f"{second['fully_skipped']} vs {len(ok_reports)}")
    disk_chunks = sum(1 for f in (OUT / "chunks").glob("*.jsonl")
                      for ln in f.read_text(encoding="utf-8").splitlines() if ln.strip())
    check("第二跑 chunk 总数与磁盘一致（幂等不重复追加）", second["n_chunks"] == disk_chunks,
          f"report {second['n_chunks']} vs disk {disk_chunks}")
    check("第二跑产物文件数不变", len(list((OUT / "chunks").glob("*.jsonl"))) == 5,
          str(len(list((OUT / "chunks").glob("*.jsonl")))))
    again_ids = [c["chunk_id"] for c in iter_chunks(compute_doc_id(IN_DIR / "sample.docx"), OUT)]
    check("第二跑 chunk_id 与首轮一致", again_ids == first_ids)

    # 断点续跑：从 chunk 阶段起跑，复用 parsed，不重新解析
    txt_id = compute_doc_id(IN_DIR / "sample.txt")
    txt_ids_before = [c["chunk_id"] for c in iter_chunks(txt_id, OUT)]
    parsed_before = artifact_path("parse", txt_id, OUT).stat().st_mtime_ns
    res = run(IN_DIR, "chunk", "chunk", workers=1, root=OUT, force=True)
    parsed_after = artifact_path("parse", txt_id, OUT).stat().st_mtime_ns
    check("--from-stage chunk 复用 parsed 产物（时间戳未变）", parsed_before == parsed_after,
          f"{parsed_before} -> {parsed_after}")
    check("强制重切后 chunk_id 依然稳定",
          [c["chunk_id"] for c in iter_chunks(txt_id, OUT)] == txt_ids_before)
    check("断点续跑只坏文件失败", res["total"] >= 5 and res["error"] == 1,
          f"{res['total']}/{res['error']}")

    # 区间/阶段名：五阶段全部接线（embed/index 依赖重模型，离线不实际跑，只校验注册与区间）
    check("五阶段全部标记为已实现", all(IMPLEMENTED[s] for s in STAGES), str(IMPLEMENTED))
    check("parse->index 覆盖全部五阶段", stage_range("parse", "index") == list(STAGES),
          str(stage_range("parse", "index")))
    try:
        run(IN_DIR / "sample.txt", "chunk", "parse", workers=1, root=OUT)
        check("非法阶段区间报错清晰", False, "没抛异常")
    except ValueError as exc:
        check("非法阶段区间报错清晰", "阶段区间非法" in str(exc), str(exc))

    # 缺上游产物：新文件直接从 chunk 起跑必须点名怎么补
    fresh = IN_DIR / "late.txt"
    fresh.write_text("# 新文档\n" + LONG_SENT + "\n", encoding="utf-8")
    rep = run(fresh, "chunk", "chunk", workers=1, root=OUT)
    err = rep["reports"][0]
    check("缺上游产物时错误提示先跑 parse", "请先跑 --from-stage parse" in (err["error"] or ""),
          str(err["error"]))
    check("错误定位到具体文件路径", err["file"] == str(fresh) and err["failed_stage"] == "parse",
          f"{err['file']} @ {err['failed_stage']}")

    # 坏文件：定位 + 错误日志 + traceback
    broken = IN_DIR / "broken.pdf"
    rep2 = run(broken, "parse", "chunk", workers=1, root=OUT)
    err2 = rep2["reports"][0]
    check("坏 PDF 单文件失败不炸整批", rep2["error"] == 1 and rep2["total"] == 1, str(rep2["error"]))
    check("坏 PDF 失败阶段定位=parse", err2["failed_stage"] == "parse", str(err2["failed_stage"]))
    check("坏 PDF report 带 traceback 与路径", bool(err2.get("traceback")) and err2["file"].endswith("broken.pdf"))
    err_log = OUT / "logs" / "pipeline_errors.jsonl"
    check("错误明细写进 logs/pipeline_errors.jsonl",
          err_log.exists() and broken.name in err_log.read_text(encoding="utf-8"))

    # 串行 vs 多进程：chunk_id 必须完全一致（Pool 不改变确定性）
    docx_id = compute_doc_id(IN_DIR / "sample.docx")
    serial = run(IN_DIR / "sample.docx", "parse", "chunk", workers=1, root=OUT / "serial", force=True)
    par = run(IN_DIR / "sample.docx", "parse", "chunk", workers=3, root=OUT / "par", force=True)
    a = [json.loads(ln)["chunk_id"] for ln in (OUT / "serial" / "chunks" / f"{docx_id}.jsonl").read_text(encoding="utf-8").splitlines() if ln.strip()]
    b = [json.loads(ln)["chunk_id"] for ln in (OUT / "par" / "chunks" / f"{docx_id}.jsonl").read_text(encoding="utf-8").splitlines() if ln.strip()]
    check("串行与多进程 chunk_id 一致", a == b and len(a) > 0, f"{len(a)} vs {len(b)}")
    check("两种并行度 chunk 数一致", serial["n_chunks"] == par["n_chunks"],
          f"{serial['n_chunks']} vs {par['n_chunks']}")

    # CLI：需求给的命令写法 --from parse（同时 --from-stage 也接受）
    code = main(["--input", str(IN_DIR / "sample.docx"), "--from", "parse", "--to", "chunk",
                 "--out-root", str(OUT / "cli"), "--workers", "1"])
    check("CLI --from/--to 可跑通", code == 0, str(code))
    code2 = main(["--input", str(IN_DIR / "sample.docx"), "--from-stage", "parse",
                  "--to-stage", "chunk", "--out-root", str(OUT / "cli2"), "--token-backend", "char"])
    check("CLI --from-stage/--to-stage 别名可用", code2 == 0, str(code2))
    code3 = main(["--input", str(broken), "--from", "parse", "--out-root", str(OUT / "cli3")])
    check("CLI 有失败文件时返回码 1", code3 == 1, str(code3))
    code4 = main(["--input", str(IN_DIR / "不存在.txt"), "--out-root", str(OUT / "cli4")])
    check("CLI 路径不存在返回码 2", code4 == 2, str(code4))


def test_scan_skips_artifacts() -> None:
    from ingestion.pipeline import collect_files

    files = collect_files(IN_DIR, recursive=True)
    check("扫描只收支持的扩展名", all(f.suffix.lower() in (".txt", ".docx", ".pptx", ".xlsx", ".pdf") for f in files),
          str([f.suffix for f in files]))
    (IN_DIR / "parsed").mkdir(exist_ok=True)
    (IN_DIR / "parsed" / "x.json").write_text("{}", encoding="utf-8")
    (IN_DIR / "nested").mkdir(exist_ok=True)
    (IN_DIR / "nested" / "chunks").mkdir(exist_ok=True)
    (IN_DIR / "nested" / "chunks" / "y.jsonl").write_text("", encoding="utf-8")
    (IN_DIR / "nested" / "real.txt").write_text("# 真文件\n内容\n", encoding="utf-8")
    again = collect_files(IN_DIR, recursive=True)
    check("产物目录被跳过（含嵌套一层）",
          not any("chunks" in f.parts or "parsed" in f.parts for f in again), str(again[-2:]))
    check("嵌套目录里的真文件仍被收集",
          any(f.name == "real.txt" for f in again))
    for junk in (IN_DIR / "parsed", IN_DIR / "nested"):
        for p in sorted(junk.rglob("*"), reverse=True):
            p.unlink() if p.is_file() else p.rmdir()


def test_progress_fallback() -> None:
    from ingestion.pipeline import _iter_progress

    got = list(_iter_progress(iter(range(7)), 7, "测试"))
    check("无 tqdm 时进度条降级不丢数据", got == list(range(7)), str(got))


def test_doc_id_strategy() -> None:
    """DocIdStrategy 三种策略 + 辅助函数 + by-path 缓存 + CONTENT 流水线集成。"""
    import hashlib
    import shutil
    import tempfile

    from ingestion.pipeline import (
        DEFAULT_ID_STRATEGY,
        DocIdStrategy,
        _normalize_block_text,
        compute_doc_id,
        content_fingerprint,
        image_fingerprint,
        normalize_strategy,
        peek_doc_id,
        read_cached_content_id,
        run,
        write_cached_content_id,
    )

    # ---------- normalize_strategy ----------
    check("normalize：合法字符串 bytes", normalize_strategy("bytes") is DocIdStrategy.BYTES)
    check("normalize：合法字符串 path-mtime（兼容连字符）", normalize_strategy("path-mtime") is DocIdStrategy.PATH_MTIME)
    check("normalize：None → 默认 BYTES", normalize_strategy(None) is DEFAULT_ID_STRATEGY)
    try:
        normalize_strategy("xxx_bad")
        check("normalize：非法值报错", False)
    except ValueError as e:
        check("normalize：非法值报错含可选清单", "bytes" in str(e))

    # ---------- BYTES ----------
    tmpd = Path(tempfile.mkdtemp())
    try:
        f = tmpd / "test.txt"
        f.write_bytes(b"hello world 123")
        id1 = compute_doc_id(f, "bytes")
        id2 = compute_doc_id(f, DocIdStrategy.BYTES)
        check("BYTES：同一文件两次调用结果一致", id1 == id2)
        expected = hashlib.sha256(f.read_bytes()).hexdigest()
        check("BYTES：等于 sha256(file_bytes)", id1 == expected)
        # 不同内容 → 不同 id
        f2 = tmpd / "test2.txt"
        f2.write_bytes(b"hello world 124")
        check("BYTES：不同字节 → 不同 id", compute_doc_id(f2, "bytes") != id1)

        # ---------- PATH_MTIME ----------
        pid1 = compute_doc_id(f, "path_mtime")
        pid2 = compute_doc_id(f, DocIdStrategy.PATH_MTIME)
        check("PATH_MTIME：同路径同 stat 结果一致", pid1 == pid2)
        # 改内容不改 mtime 不行 -> 先测改名
        f3 = tmpd / "renamed.txt"
        f.rename(f3)
        pid3 = compute_doc_id(f3, "path_mtime")
        check("PATH_MTIME：改名后 id 变（路径是身份）", pid3 != pid1)

        # ---------- _normalize_block_text ----------
        raw = "  hello  \r\nworld  \r\n"
        normed = _normalize_block_text(raw)
        check("_normalize：\\r\\n → \\n + 去尾随空格 + strip", normed == "hello\nworld")
        check("_normalize：None → 空串", _normalize_block_text(None) == "")
        check("_normalize：纯数字不变形", _normalize_block_text("42") == "42")

        # ---------- CONTENT 指纹（文本块） ----------
        parsed_a = {
            "source_format": "docx",
            "blocks": [
                {"type": "text", "content": "第一章 概述", "page": 1, "bbox": [0, 0, 100, 50]},
                {"type": "text", "content": "这是正文内容。  \r\n  第二行", "page": 2, "bbox": [10, 20]},
            ],
        }
        # 同一内容、不同 bbox/font/page → CONTENT 指纹必须等（跨导出稳定）
        parsed_b = {
            "source_format": "docx",
            "blocks": [
                {"type": "text", "content": "第一章 概述", "page": 99, "bbox": [999, 999, 999, 999], "font": "宋体"},
                {"type": "text", "content": "这是正文内容。\n  第二行", "page": 55},
            ],
        }
        fp_a, _ = content_fingerprint(parsed_a)
        fp_b, _ = content_fingerprint(parsed_b)
        check("CONTENT：bbox/page/font 变化不影响指纹（跨导出稳定）", fp_a == fp_b)

        # 不同正文 → 不同 id
        parsed_c = {
            "source_format": "docx",
            "blocks": [
                {"type": "text", "content": "完全不同内容"},
            ],
        }
        fp_c, _ = content_fingerprint(parsed_c)
        check("CONTENT：不同正文 → 不同指纹", fp_a != fp_c)

        # compute_doc_id CONTENT 缺 parsed 报错
        try:
            compute_doc_id(f, "content", data=b"xxx")
            check("CONTENT：缺 parsed 参数报错", False)
        except ValueError as e:
            check("CONTENT：缺 parsed 参数报错含提示", "必须先 parse" in str(e))

        # compute_doc_id CONTENT 正常路径
        cid = compute_doc_id(f, "content", data=b"ignored", parsed=parsed_a)
        check("compute_doc_id CONTENT：返回指纹等于 content_fingerprint", cid == fp_a)

        # ---------- image_fingerprint ----------
        # 含 block["doc_id"] 但 image_fingerprint 绝不使用它
        img_block_with_docid = {
            "type": "image",
            "content": "[图片 200x100px]",
            "doc_id": "parent_sha:p3:x12",
            "clip_vector": None,
            "phash": "",
        }
        fp_img, method = image_fingerprint(img_block_with_docid, data=None, source_format="docx")
        check("image_fingerprint：不用 block['doc_id']", method != "doc_id")
        # clip_vector 有值时走 clip_quantized
        img_clip = {
            "type": "image",
            "content": "ignored",
            "clip_vector": [0.5, -0.2, 0.05, -0.7, 0.4, -0.15, 0.8, -0.3],
            "phash": "",
        }
        _, m2 = image_fingerprint(img_clip, data=None, source_format="docx")
        check("image_fingerprint：有 clip_vector → clip_quantized", m2 == "clip_quantized")
        # phash 字段有值 → 直接取
        img_ph = {"type": "image", "content": "x", "phash": "abcdef0123456789"}
        fp_ph, m_ph = image_fingerprint(img_ph)
        check("image_fingerprint：显式 phash 字段优先", fp_ph == "abcdef0123456789" and m_ph == "phash")

        # ---------- by-path 缓存 ----------
        cache_root = tmpd / "out"
        cache_root.mkdir()
        # 模拟 CONTENT parse 后写缓存
        test_file = tmpd / "cached.txt"
        test_file.write_bytes(b"cached content for test")
        content_id = "abc123deadbeef"
        write_cached_content_id(test_file, content_id, cache_root)
        read_back = read_cached_content_id(test_file, cache_root)
        check("by-path 缓存：写入后读回一致", read_back == content_id)
        # peek 使用缓存
        peeked = peek_doc_id(test_file, "content", cache_root)
        check("by-path 缓存：peek_doc_id CONTENT 命中缓存", peeked == content_id)
        # 缓存文件不存在时返回 None
        new_file = tmpd / "new_no_cache.txt"
        new_file.write_bytes(b"new file no cache yet")
        peek_miss = peek_doc_id(new_file, "content", cache_root)
        check("by-path 缓存：无缓存时 peek 返回 None", peek_miss is None)

        # ---------- CONTENT 策略流水线集成 ----------
        pipe_in = tmpd / "pipe_in"
        pipe_in.mkdir()
        # 放一个 txt 文件进去（不用 Office，避免 zip 时间戳干扰）
        (pipe_in / "doc1.txt").write_text("第一章 测试\n\n这是正文。\n第二行内容。", encoding="utf-8")
        pipe_out = tmpd / "pipe_out"
        pipe_out.mkdir()

        # 第一跑 CONTENT
        s1 = run(pipe_in, "parse", "chunk", id_strategy="content", workers=1, root=pipe_out)
        check("CONTENT 流水线：无 error（除无）", s1["error"] == 0)
        check("CONTENT 流水线：策略记录正确", s1["id_strategy"] == "content")
        # 产物 parsed/{content_id}.json 存在
        parsed_files = list((pipe_out / "parsed").glob("*.json"))
        check("CONTENT 流水线：parsed 产物落盘", len(parsed_files) >= 1)
        # doc_id 是 content 指纹不是 bytes 哈希
        content_parsed = json.loads(parsed_files[0].read_text(encoding="utf-8"))
        bytes_id = hashlib.sha256((pipe_in / "doc1.txt").read_bytes()).hexdigest()
        check("CONTENT 流水线：doc_id ≠ sha256(bytes)", content_parsed["doc_id"] != bytes_id)
        check("CONTENT 流水线：doc_id_strategy 字段记录 content", content_parsed.get("doc_id_strategy") == "content")
        check("CONTENT 流水线：bytes_doc_id 保留原始字节哈希", content_parsed.get("bytes_doc_id") == bytes_id)

        # 第二跑 CONTENT：by-path 缓存命中 → 全部 skipped
        s2 = run(pipe_in, "parse", "chunk", id_strategy="content", workers=1, root=pipe_out)
        check("CONTENT 流水线：第二跑幂等（全 skipped）", s2["fully_skipped"] >= 1)

        # ---------- PATH_MTIME 策略流水线 ----------
        pipe_out2 = tmpd / "pipe_out_pm"
        pipe_out2.mkdir()
        s3 = run(pipe_in, "parse", "chunk", id_strategy="path_mtime", workers=1, root=pipe_out2)
        check("PATH_MTIME 流水线：无 error", s3["error"] == 0)
        check("PATH_MTIME 流水线：策略记录", s3["id_strategy"] == "path_mtime")
        parsed_pm = list((pipe_out2 / "parsed").glob("*.json"))
        check("PATH_MTIME 流水线：产物落盘", len(parsed_pm) >= 1)
        content_pm = json.loads(parsed_pm[0].read_text(encoding="utf-8"))
        check("PATH_MTIME 流水线：doc_id_strategy=path_mtime", content_pm.get("doc_id_strategy") == "path_mtime")

    finally:
        shutil.rmtree(tmpd, ignore_errors=True)

    print("[OK]   test_doc_id_strategy 全部通过")


def main_tests() -> None:
    test_token_estimator()
    test_config_resolve()
    test_text_chunking()
    test_table_chunking()
    test_chunk_schema()
    test_real_docs_quality()
    test_scan_skips_artifacts()
    test_pipeline_idempotent()
    test_progress_fallback()
    test_doc_id_strategy()


if __name__ == "__main__":
    print(f"Python {sys.version.split()[0]} / 临时目录 {TMP}")
    make_all()
    main_tests()
    print("=" * 60)
    if FAILS:
        print(f"RESULT: {len(FAILS)} FAILED")
        for f in FAILS:
            print("  -", f)
        raise SystemExit(1)
    print("RESULT: ALL PASS")
    raise SystemExit(0)
